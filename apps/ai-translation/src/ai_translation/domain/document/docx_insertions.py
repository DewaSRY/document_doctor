"""Writes the blocks a user added in the editor (paragraphs, lists, tables, images) into a DOCX."""

from io import BytesIO

from docx.document import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_HEADING_KINDS = {"normal", "title", "subtitle", "h1", "h2", "h3", "h4", "h5", "h6"}
# The editor draws the same markers and colours (insert-extension.ts in the portal).
_BULLETS = ["•", "◦", "▪", "•", "◦"]
_NUMBER_FORMATS = ["decimal", "lowerLetter", "lowerRoman", "decimal", "lowerLetter"]
_TODO = {False: "☐ ", True: "☒ "}
_QUOTE_COLOR = "595959"
_CODE_FONT = "Courier New"
_CODE_SHADING = "F2F2F2"
_HEADER_SHADING = "F2F2F2"
_RULE_COLOR = "BFBFBF"
# Levels of a list, as in Word's default lists: half an inch each, the marker hanging a quarter inch.
_LIST_INDENT = 720
_LIST_HANGING = 360


class InsertionWriter:
    """
    Places inserted blocks after elements of the body, by the index each one
    had in the original upload. Run it after the translations are written:
    it adds paragraphs, which would change the keys of the segments after them.
    """

    def __init__(self, doc: Document, media: dict[str, bytes], target_language: str):
        from .docx_handler import DOCXHandler

        self.doc = doc
        self.media = media
        self.target_language = target_language
        self._handler = DOCXHandler
        self._paragraph_styles = {style.style_id for style in doc.styles if style.type == 1}
        self._bullet_num: str | None = None
        self._number_abstract: str | None = None
        section = doc.sections[-1]
        self._text_width = Emu(section.page_width - section.left_margin - section.right_margin)

    def write(self, blocks: list[dict]) -> None:
        body = self.doc.element.body
        # Indexes refer to the original body, which rebuilding does not change: capture it before adding to it.
        elements = list(body.iterchildren())
        content = [element for element in elements if element.tag != qn("w:sectPr")]
        if not content:
            return

        # A numbered list is one run of numbered blocks with nothing between them; each starts at 1.
        previous: dict | None = None
        list_num: str | None = None

        # Blocks after the same element go in order after it.
        cursor = None
        cursor_after = None
        for block in blocks:
            after = block.get("after", -1)
            if block.get("type") == "paragraph" and block.get("kind") == "numbered":
                continues = (
                    previous is not None
                    and previous.get("type") == "paragraph"
                    and previous.get("kind") == "numbered"
                    and previous.get("after") == after
                )
                if not continues:
                    list_num = self._new_number_list()
            previous = block

            new = self._element(block, list_num)
            if new is None:
                continue
            if after != cursor_after or cursor is None:
                cursor_after = after
                if after < 0:
                    content[0].addprevious(new)
                    cursor = new
                    continue
                anchor = elements[after] if after < len(elements) else content[-1]
                if anchor.tag == qn("w:sectPr"):
                    anchor.addprevious(new)
                else:
                    anchor.addnext(new)
            else:
                cursor.addnext(new)
            cursor = new

    # ------------------------------------------------------------------ blocks

    def _element(self, block: dict, list_num: str | None):
        kind = block.get("type")
        if kind == "paragraph":
            return self._paragraph(block, list_num)
        if kind == "table":
            return self._table(block)
        if kind == "image":
            return self._image(block)
        if kind == "divider":
            paragraph = self._new_paragraph()
            ppr = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for name, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", _RULE_COLOR)):
                bottom.set(qn(name), value)
            borders.append(bottom)
            ppr.insert_element_before(
                borders, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr"
            )
            paragraph.paragraph_format.space_after = Pt(6)
            return paragraph._p
        if kind == "page_break":
            paragraph = self._new_paragraph()
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            paragraph.add_run()._r.append(br)
            return paragraph._p
        return None

    def _new_paragraph(self) -> Paragraph:
        return Paragraph(OxmlElement("w:p"), self.doc._body)

    def _paragraph(self, block: dict, list_num: str | None):
        paragraph = self._new_paragraph()
        kind = block.get("kind", "normal")
        level = min(4, max(0, int(block.get("level") or 0)))
        ppr = paragraph._p.get_or_add_pPr()
        fmt = paragraph.paragraph_format
        fallback: tuple[float, bool] | None = None

        if kind in _HEADING_KINDS and kind != "normal":
            from .docx_handler import _HEADING_FALLBACKS, _HEADING_STYLE_IDS

            style_id = _HEADING_STYLE_IDS[kind]
            if style_id in self._paragraph_styles:
                ppr.get_or_add_pStyle().val = style_id
            else:
                fallback = _HEADING_FALLBACKS[kind]
        elif kind == "quote":
            # Drawn the same in the editor: indented, a grey rule on the left, italic grey text.
            borders = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            for name, value in (("w:val", "single"), ("w:sz", "18"), ("w:space", "8"), ("w:color", _RULE_COLOR)):
                left.set(qn(name), value)
            borders.append(left)
            ppr.insert_element_before(
                borders, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr"
            )
            fmt.left_indent = Pt(12)
        elif kind == "code":
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:color"), "auto")
            shading.set(qn("w:fill"), _CODE_SHADING)
            ppr.insert_element_before(
                shading, "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr"
            )
            fmt.space_after = Pt(0)
        elif kind == "bullet":
            self._set_numbering(paragraph, self._bullet_list(), level)
        elif kind == "numbered" and list_num is not None:
            self._set_numbering(paragraph, list_num, level)
        elif kind == "todo":
            fmt.left_indent = Pt(18 * level)

        if block.get("align") in _ALIGNMENTS:
            paragraph.alignment = _ALIGNMENTS[block["align"]]

        if kind == "todo":
            self._add_run(paragraph, _TODO[bool(block.get("checked"))], {})
        # Runs carry only the toggles the user turned on, so the paragraph's style shows through.
        for run in block.get("runs") or []:
            style = run.get("style") or {}
            new_run = self._add_run(paragraph, run.get("text", ""), style)
            if new_run is None:
                continue
            if fallback:
                new_run.font.size = Pt(fallback[0])
                if style.get("bold") is None:
                    new_run.font.bold = fallback[1]
            if kind == "quote":
                if style.get("italic") is None:
                    new_run.font.italic = True
                if not style.get("color"):
                    new_run.font.color.rgb = RGBColor.from_string(_QUOTE_COLOR)
            if kind == "code":
                self._set_font(new_run, _CODE_FONT)
        return paragraph._p

    def _add_run(self, paragraph: Paragraph, text: str, style: dict) -> Run | None:
        if not text:
            return None
        run = paragraph.add_run()
        # Line breaks the user typed (Shift+Enter) become <w:br/>.
        run.text = text
        self._handler._apply_run_style(run, style)
        self._handler._set_language(run, self.target_language)
        return run

    @staticmethod
    def _set_font(run: Run, family: str) -> None:
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        for name in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(name), family)

    def _table(self, block: dict):
        rows = block.get("rows") or []
        columns = max((len(row.get("cells") or []) for row in rows), default=0)
        if not rows or not columns:
            return None
        # Built at the end of the body, then moved into place by write().
        table = self.doc.add_table(rows=len(rows), cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        grid = next((style for style in self.doc.styles if style.style_id == "TableGrid"), None)
        if grid is not None:
            table.style = grid
        else:
            self._set_grid_borders(table._tbl)
        width = Emu(int(self._text_width / columns))
        for column in table.columns:
            column.width = width

        header = bool(block.get("header_row"))
        for row_index, row in enumerate(rows):
            cells = row.get("cells") or []
            if header and row_index == 0:
                tr_pr = table.rows[0]._tr.get_or_add_trPr()
                tr_pr.append(OxmlElement("w:tblHeader"))
            for col_index in range(columns):
                cell = table.cell(row_index, col_index)
                cell.width = width
                paragraphs = cells[col_index].get("paragraphs") if col_index < len(cells) else None
                for index, runs in enumerate(paragraphs or [[]]):
                    paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in runs:
                        style = dict(run.get("style") or {})
                        if header and row_index == 0 and style.get("bold") is None:
                            style["bold"] = True
                        self._add_run(paragraph, run.get("text", ""), style)
                if header and row_index == 0:
                    tc_pr = cell._tc.get_or_add_tcPr()
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:val"), "clear")
                    shading.set(qn("w:color"), "auto")
                    shading.set(qn("w:fill"), _HEADER_SHADING)
                    tc_pr.append(shading)
        return table._tbl

    @staticmethod
    def _set_grid_borders(tbl) -> None:
        """Thin borders on every edge, like Word's "Table Grid", when the document lacks that style."""
        tbl_pr = tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = OxmlElement(f"w:{side}")
            for name, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "auto")):
                edge.set(qn(name), value)
            borders.append(edge)
        tbl_pr.insert_element_before(
            borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"
        )

    def _image(self, block: dict):
        data = self.media.get(block.get("src", ""))
        if data is None:
            return None
        paragraph = self._new_paragraph()
        paragraph.alignment = _ALIGNMENTS.get(block.get("align", "center"), WD_ALIGN_PARAGRAPH.CENTER)
        run = paragraph.add_run()
        run.add_picture(BytesIO(data), width=Pt(block["width"]), height=Pt(block["height"]))
        alt = block.get("alt")
        if alt:
            for doc_pr in run._r.iter(qn("wp:docPr")):
                doc_pr.set("descr", alt)
        return paragraph._p

    # ------------------------------------------------------------------ lists

    def _set_numbering(self, paragraph: Paragraph, num_id: str, level: int) -> None:
        ppr = paragraph._p.get_or_add_pPr()
        num_pr = ppr.get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = level
        num_pr.get_or_add_numId().val = int(num_id)

    def _bullet_list(self) -> str:
        if self._bullet_num is None:
            abstract = self._add_abstract(bullet=True)
            self._bullet_num = self._add_num(abstract, restart=False)
        return self._bullet_num

    def _new_number_list(self) -> str:
        if self._number_abstract is None:
            self._number_abstract = self._add_abstract(bullet=False)
        return self._add_num(self._number_abstract, restart=True)

    def _numbering(self):
        return self.doc.part.numbering_part.element

    def _add_abstract(self, bullet: bool) -> str:
        numbering = self._numbering()
        ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.iterfind(qn("w:abstractNum"))]
        abstract_id = str(max(ids, default=-1) + 1)
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), abstract_id)
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "hybridMultilevel")
        abstract.append(multi)
        for level in range(5):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            for tag, value in (
                ("w:start", "1"),
                ("w:numFmt", "bullet" if bullet else _NUMBER_FORMATS[level]),
                ("w:lvlText", _BULLETS[level] if bullet else f"%{level + 1}."),
                ("w:lvlJc", "left"),
            ):
                child = OxmlElement(tag)
                child.set(qn("w:val"), value)
                lvl.append(child)
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(_LIST_INDENT * (level + 1)))
            ind.set(qn("w:hanging"), str(_LIST_HANGING))
            ppr.append(ind)
            lvl.append(ppr)
            abstract.append(lvl)
        # Every abstractNum comes before the first num.
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        return abstract_id

    def _add_num(self, abstract_id: str, restart: bool) -> str:
        numbering = self._numbering()
        ids = [int(n.get(qn("w:numId"))) for n in numbering.iterfind(qn("w:num"))]
        num_id = str(max(ids, default=0) + 1)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), num_id)
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_id)
        num.append(ref)
        if restart:
            # Numbers of different num elements of one abstractNum would otherwise continue.
            for level in range(5):
                override = OxmlElement("w:lvlOverride")
                override.set(qn("w:ilvl"), str(level))
                start = OxmlElement("w:startOverride")
                start.set(qn("w:val"), "1")
                override.append(start)
                num.append(override)
        cleanup = numbering.find(qn("w:numIdMacAtCleanup"))
        if cleanup is not None:
            cleanup.addprevious(num)
        else:
            numbering.append(num)
        return num_id
