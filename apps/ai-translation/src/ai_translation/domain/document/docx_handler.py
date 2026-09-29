import re
from collections import Counter
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO

from docx import Document as DocxDocument
from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from .base import DocumentHandler
from .docx_insertions import InsertionWriter
from .docx_layout import DocxLayoutReader

# Tabs and line breaks split a paragraph into separately translated pieces,
# so "Label<tab>: value" and multi-line addresses keep their structure.
_SEPARATORS = {qn("w:tab"), qn("w:ptab"), qn("w:br"), qn("w:cr")}
_LINE_BREAKS = {qn("w:br"), qn("w:cr")}
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
# Language tags, so Word proofs (and picks fonts for) the translation in the target language.
_LANGUAGE_TAGS = {
    "en": "en-US",
    "id": "id-ID",
    "fr": "fr-FR",
    "de": "de-DE",
    "es": "es-ES",
    "ja": "ja-JP",
    "ko": "ko-KR",
    "zh": "zh-CN",
}
_EAST_ASIAN = {"ja", "ko", "zh"}
_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
# Word's built-in style ids for the editor's paragraph styles.
_HEADING_STYLE_IDS = {
    "title": "Title",
    "subtitle": "Subtitle",
    **{f"h{level}": f"Heading{level}" for level in range(1, 7)},
}
# Direct formatting for a heading when the document has no such style; the
# editor shows the same (HEADING_FALLBACKS in the portal).
_HEADING_FALLBACKS = {
    "title": (26.0, False),
    "subtitle": (15.0, False),
    "h1": (20.0, True),
    "h2": (16.0, True),
    "h3": (14.0, True),
    "h4": (12.0, True),
    "h5": (11.0, True),
    "h6": (11.0, False),
}
# Elements that follow w:shd in a run's properties (the schema's order).
_AFTER_SHD = (
    "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
    "w:eastAsianLayout", "w:specVanish", "w:oMath",
)


@dataclass
class _Piece:
    """Text of a paragraph between tabs and line breaks: the unit that is translated."""

    paragraph: Paragraph
    items: list[Run | Hyperlink]

    @property
    def text(self) -> str:
        return "".join(item.text for item in self.items)


class _PartParent:
    """Minimal parent for paragraphs found outside python-docx's object model."""

    def __init__(self, part):
        self.part = part


class DOCXHandler(DocumentHandler):
    """
    Handler for DOCX documents.

    Every paragraph (body, tables, headers, footers, text boxes and content
    controls) is one segment; a paragraph with tabs or line breaks is one
    segment per piece between them. Keys look like 'para_3' or 'para_3_1'
    and are stable between extraction and rebuild.
    """

    _HAS_LETTER = re.compile(r"[^\W\d_]")

    async def extract_text(self, file_content: bytes) -> dict[str, list[str]]:
        """Extract translatable paragraphs from the DOCX."""
        doc = DocxDocument(BytesIO(file_content))
        return {
            key: self._segment_text(piece.text)
            for key, piece in self._collect_pieces(doc).items()
        }

    async def create_translated_document(
        self,
        file_content: bytes,
        translations: dict[str, str],
        source_language: str,
        target_language: str,
        styles: dict[str, dict] | None = None,
        runs: dict[str, list[dict]] | None = None,
        insertions: list[dict] | None = None,
        media: dict[str, bytes] | None = None,
    ) -> bytes:
        """
        Replace each paragraph with its translation, keeping the formatting.

        styles: a user's overrides by key: the font and size of the segment's
        text, and the paragraph style, alignment, spacing and indent of its paragraph.
        runs: the formatting of ranges of a segment's text (see _apply_runs).
        insertions: blocks added in the editor, each after an element of the
        original body (see InsertionWriter); media: their images, by name.
        """
        styles = styles or {}
        runs = runs or {}
        doc = DocxDocument(BytesIO(file_content))
        styled = set()

        for key, piece in self._collect_pieces(doc).items():
            translated_text = translations.get(key, "").strip()
            if not translated_text:
                continue
            style = styles.get(key) or {}
            # Paragraph formatting is sent with every piece of the paragraph; apply it once.
            if style and piece.paragraph._p not in styled:
                styled.add(piece.paragraph._p)
                self._apply_paragraph_style(doc, piece.paragraph, style)
            self._write_piece(piece, translated_text, target_language)
            self._apply_text_style(piece, style)
            if runs.get(key):
                self._apply_runs(piece, runs[key], translations[key])
            self._let_row_grow(piece.paragraph)

        for element, _ in self._stories(doc):
            self._sync_fallback_text_boxes(element)

        # Last: new paragraphs would change the keys of the segments after them.
        if insertions:
            InsertionWriter(doc, media or {}, target_language).write(insertions)

        doc.core_properties.subject = (
            f"Translated from {source_language} to {target_language}"
        )
        doc.core_properties.comments = "Translated by AI Translation Service"
        doc.core_properties.language = _LANGUAGE_TAGS.get(
            target_language, target_language
        )

        output = BytesIO()
        doc.save(output)
        return output.getvalue()

    def document_layout(self, file_content: bytes) -> dict:
        """Pages, paragraphs, tables and images of the document, for the editor."""
        doc = DocxDocument(BytesIO(file_content))
        return DocxLayoutReader(doc, self._collect_pieces(doc)).read()

    @staticmethod
    def media(file_content: bytes, name: str) -> tuple[bytes, str] | None:
        """An image of the document by its part name (as the layout returns it)."""
        doc = DocxDocument(BytesIO(file_content))
        for part in doc.part.package.iter_parts():
            if str(part.partname).lstrip("/") == name and name.startswith("word/media/"):
                return part.blob, part.content_type
        return None

    # ------------------------------------------------------------------ extraction

    def _collect_pieces(self, doc: Document) -> dict[str, _Piece]:
        result: dict[str, _Piece] = {}
        for number, paragraph in enumerate(self._collect_paragraphs(doc)):
            pieces = self._split_paragraph(paragraph)
            if len(pieces) == 1:
                result[f"para_{number}"] = pieces[0]
                continue
            for index, piece in enumerate(pieces):
                if self._HAS_LETTER.search(piece.text):
                    result[f"para_{number}_{index}"] = piece
        return result

    def _collect_paragraphs(self, doc: Document) -> list[Paragraph]:
        paragraphs: list[Paragraph] = []

        def walk(container) -> None:
            """Add the paragraphs of a document, cell, header or footer, including nested tables."""
            paragraphs.extend(container.paragraphs)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        walk(cell)

        walk(doc)
        for section in doc.sections:
            # Linked headers belong to an earlier section (or do not exist); reading them would add one.
            for header_footer in (section.header, section.footer):
                if not header_footer.is_linked_to_previous:
                    walk(header_footer)

        # Paragraphs python-docx does not expose: text boxes, content controls and
        # first-page / even-page headers. They come last, so keys of the rest stay the same.
        for element, parent in self._stories(doc):
            for p in element.iter(qn("w:p")):
                # The fallback copy of a text box is synced after writing.
                if next(p.iterancestors(f"{_MC}Fallback"), None) is None:
                    paragraphs.append(Paragraph(p, parent))

        # Merged cells and linked headers return the same paragraph more than once.
        result: list[Paragraph] = []
        seen = set()
        for paragraph in paragraphs:
            if paragraph._p in seen:
                continue
            seen.add(paragraph._p)
            if self._HAS_LETTER.search(paragraph.text):
                result.append(paragraph)
        return result

    @staticmethod
    def _stories(doc: Document) -> list[tuple]:
        """The body and every header and footer part, each with a parent for its paragraphs."""
        stories = [(doc.element.body, doc._body)]
        for rel in doc.part.rels.values():
            if rel.reltype in (RT.HEADER, RT.FOOTER) and not rel.is_external:
                stories.append((rel.target_part.element, _PartParent(rel.target_part)))
        return stories

    def _split_paragraph(self, paragraph: Paragraph) -> list[_Piece]:
        self._isolate_separators(paragraph)
        pieces = [_Piece(paragraph, [])]
        separators: list[Run] = []
        for item in paragraph.iter_inner_content():
            if isinstance(item, Run) and self._separator(item) is not None:
                separators.append(item)
                pieces.append(_Piece(paragraph, []))
            else:
                pieces[-1].items.append(item)

        # A line break inside a sentence (manual wrapping) does not split it.
        merged = [pieces[0]]
        for separator, piece in zip(separators, pieces[1:], strict=True):
            previous = merged[-1]
            tag = self._separator(separator)
            continues = (
                tag.tag in _LINE_BREAKS
                and tag.get(qn("w:type")) in (None, "textWrapping")
                and previous.text.strip()
                and not re.search(r"[.:;!?。：]\s*$", previous.text)
                and piece.text.lstrip()[:1].islower()
            )
            if continues:
                previous.items.extend([separator, *piece.items])
            else:
                merged.append(piece)
        return merged

    @staticmethod
    def _separator(run: Run):
        """The tab or break element if the run holds nothing else."""
        children = [child for child in run._r if child.tag != qn("w:rPr")]
        if len(children) == 1 and children[0].tag in _SEPARATORS:
            return children[0]
        return None

    @staticmethod
    def _isolate_separators(paragraph: Paragraph) -> None:
        """Move tabs and breaks inside a run with text into runs of their own."""
        for r in paragraph._p.xpath("./w:r"):
            children = [child for child in r if child.tag != qn("w:rPr")]
            if len(children) < 2 or not any(
                child.tag in _SEPARATORS for child in children
            ):
                continue
            groups: list[list] = []
            for child in children:
                if (
                    child.tag in _SEPARATORS
                    or not groups
                    or groups[-1][-1].tag in _SEPARATORS
                ):
                    groups.append([child])
                else:
                    groups[-1].append(child)
            for group in groups:
                new_r = OxmlElement("w:r", attrs=dict(r.attrib))
                if r.rPr is not None:
                    new_r.append(deepcopy(r.rPr))
                new_r.extend(group)
                r.addprevious(new_r)
            r.getparent().remove(r)

    # ------------------------------------------------------------------ rebuild

    @classmethod
    def _write_piece(
        cls, piece: _Piece, translated_text: str, target_language: str
    ) -> None:
        """
        Write the translation back with the formatting that covers most of the text.

        Hyperlinks whose text still appears in the translation (URLs, e-mail
        addresses) are left untouched, so the link keeps working; the translated
        text around them goes into the plain runs before and after each link.
        """
        items = piece.items
        paragraph = piece.paragraph

        kept_links: list[tuple[int, Hyperlink]] = []
        pieces: list[str] = []
        cursor = 0
        for index, item in enumerate(items):
            if not isinstance(item, Hyperlink) or not item.text.strip():
                continue
            position = translated_text.find(item.text, cursor)
            if position < 0:
                continue
            kept_links.append((index, item))
            pieces.append(translated_text[cursor:position])
            cursor = position + len(item.text)
        pieces.append(translated_text[cursor:])

        all_runs = [
            run
            for item in items
            if not any(item is link for _, link in kept_links)
            for run in (item.runs if isinstance(item, Hyperlink) else [item])
            if run.text
        ]
        template = cls._dominant_run(all_runs)
        bounds = [-1] + [index for index, _ in kept_links] + [len(items)]

        for group, text in enumerate(pieces):
            group_items = items[bounds[group] + 1 : bounds[group + 1]]
            runs: list[Run] = []
            for item in group_items:
                runs.extend(item.runs if isinstance(item, Hyperlink) else [item])

            # Only touch runs with text, so images and other inline objects survive.
            text_runs = [run for run in runs if run.text]
            if text_runs:
                written = cls._fill_runs(text_runs, text, template)
            elif text.strip():
                new_run = Run(
                    deepcopy(template._r) if template else paragraph._p.add_r(),
                    paragraph,
                )
                new_run.text = text
                if group > 0:
                    kept_links[group - 1][1]._hyperlink.addnext(new_run._r)
                elif kept_links:
                    kept_links[0][1]._hyperlink.addprevious(new_run._r)
                written = [new_run]
            else:
                continue

            for run in written:
                cls._set_language(run, target_language)

    @staticmethod
    def _text_runs(piece: _Piece) -> list[Run]:
        """The runs that hold the piece's text, in order."""
        return [
            run
            for item in piece.items
            for run in (item.runs if isinstance(item, Hyperlink) else [item])
            if run.text
        ]

    @classmethod
    def _apply_paragraph_style(cls, doc: Document, paragraph: Paragraph, style: dict) -> None:
        heading = style.get("heading")
        if heading:
            style_id = _HEADING_STYLE_IDS.get(heading)
            if heading == "normal":
                style_id = next(
                    (
                        s.style_id
                        for s in doc.styles
                        if s.type == 1 and getattr(s, "element", None) is not None
                        and s.element.get(qn("w:default")) in ("1", "true")
                    ),
                    None,
                )
            known = style_id is not None and any(
                s.style_id == style_id for s in doc.styles if s.type == 1
            )
            ppr = paragraph._p.get_or_add_pPr()
            if known:
                ppr.get_or_add_pStyle().val = style_id
            elif heading == "normal":
                ppr._remove_pStyle()
            # Text formatting of the old style would override the new one.
            for run in paragraph.runs:
                rpr = run._r.rPr
                if rpr is None:
                    continue
                for tag in ("w:b", "w:bCs", "w:i", "w:iCs", "w:sz", "w:szCs", "w:color", "w:rFonts"):
                    for element in rpr.findall(qn(tag)):
                        rpr.remove(element)
            if not known and heading in _HEADING_FALLBACKS:
                size, bold = _HEADING_FALLBACKS[heading]
                for run in paragraph.runs:
                    run.font.size = Pt(size)
                    run.font.bold = bold

        fmt = paragraph.paragraph_format
        if style.get("align") in _ALIGNMENTS:
            paragraph.alignment = _ALIGNMENTS[style["align"]]
        if style.get("line_spacing") is not None:
            fmt.line_spacing = float(style["line_spacing"])
        if style.get("space_before") is not None:
            fmt.space_before = Pt(style["space_before"])
        if style.get("space_after") is not None:
            fmt.space_after = Pt(style["space_after"])
        if style.get("indent_left") is not None:
            fmt.left_indent = Pt(style["indent_left"])

    @classmethod
    def _apply_text_style(cls, piece: _Piece, style: dict) -> None:
        """A font or size chosen for the whole segment."""
        if style.get("font_size") is None and not style.get("family"):
            return
        for run in cls._text_runs(piece):
            if style.get("font_size") is not None:
                run.font.size = Pt(style["font_size"])
                rpr = run._r.get_or_add_rPr()
                for element in rpr.findall(qn("w:szCs")):
                    rpr.remove(element)
            if style.get("family"):
                fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
                for name in ("w:ascii", "w:hAnsi", "w:cs"):
                    fonts.set(qn(name), style["family"])
                for name in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):
                    fonts.attrib.pop(qn(name), None)

    @classmethod
    def _apply_runs(cls, piece: _Piece, runs: list[dict], text: str) -> None:
        """
        Give ranges of the written translation their own formatting ("a **bold**
        word"): every run holding the text is split where the ranges change, and
        each part keeps the run's formatting plus its range's style.

        runs: [{"text": str, "style": {bold, italic, underline, strike, color, highlight}}],
        whose texts join to the (unstripped) translation.
        """
        start = -(len(text) - len(text.lstrip()))
        ranges = []
        for run in runs:
            end = start + len(run.get("text", ""))
            ranges.append((start, end, run.get("style") or {}))
            start = end

        offset = 0
        for run in cls._text_runs(piece):
            value = run.text
            begin, finish = offset, offset + len(value)
            offset = finish
            parts = [
                (max(a, begin), min(b, finish), style)
                for a, b, style in ranges
                if min(b, finish) > max(a, begin)
            ]
            if not parts:
                continue
            anchor = run._r
            for a, b, style in parts:
                new_run = Run(deepcopy(run._r), run._parent)
                new_run.text = value[a - begin : b - begin]
                cls._apply_run_style(new_run, style)
                anchor.addnext(new_run._r)
                anchor = new_run._r
            run._r.getparent().remove(run._r)

    @staticmethod
    def _apply_run_style(run: Run, style: dict) -> None:
        font = run.font
        for name in ("bold", "italic", "strike"):
            if style.get(name) is not None:
                setattr(font, name, bool(style[name]))
        if style.get("underline") is not None:
            font.underline = bool(style["underline"])
        if style.get("color"):
            font.color.rgb = RGBColor.from_string(style["color"].lstrip("#").upper())
        highlight = style.get("highlight")
        if highlight:
            rpr = run._r.get_or_add_rPr()
            rpr._remove_highlight()
            for element in rpr.findall(qn("w:shd")):
                rpr.remove(element)
            if highlight != "transparent":
                # Any colour, not only Word's named highlight colours.
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), highlight.lstrip("#").upper())
                rpr.insert_element_before(shd, *_AFTER_SHD)

    @classmethod
    def _fill_runs(cls, runs: list[Run], text: str, template: Run | None) -> list[Run]:
        """
        Put the text into the first run and empty the others.

        A lead-in with its own formatting ("**Note:** text") keeps it, when the
        translation has a colon as well; the rest gets the main formatting.
        """
        lead = []
        for run in runs:
            if lead and cls._format_key(run) != cls._format_key(lead[0]):
                break
            lead.append(run)
        lead_text = "".join(run.text for run in lead)
        split = re.search(r"[:：]", text)
        if (
            template is not None
            and len(lead) < len(runs)
            and cls._format_key(lead[0]) != cls._format_key(template)
            and re.search(r"[:：]\s*$", lead_text)
            and split
        ):
            head, rest = text[: split.end()], text[split.end() :]
            first_rest = runs[len(lead)]
            for run in runs:
                run.text = ""
            lead[0].text = head
            first_rest.text = rest
            cls._copy_format(template, first_rest)
            return [lead[0], first_rest]

        for run in runs:
            run.text = ""
        runs[0].text = text
        if template is not None:
            cls._copy_format(template, runs[0])
        return [runs[0]]

    @classmethod
    def _dominant_run(cls, runs: list[Run]) -> Run | None:
        """The run whose formatting covers the most characters."""
        if not runs:
            return None
        weights: Counter[bytes] = Counter()
        first: dict[bytes, Run] = {}
        for run in runs:
            key = cls._format_key(run)
            weights[key] += len(run.text.strip())
            first.setdefault(key, run)
        return first[max(weights, key=lambda key: weights[key])]

    @staticmethod
    def _format_key(run: Run) -> bytes:
        rPr = run._r.rPr
        return etree.tostring(rPr) if rPr is not None else b""

    @staticmethod
    def _copy_format(source: Run, target: Run) -> None:
        if source._r is target._r:
            return
        if target._r.rPr is not None:
            target._r.remove(target._r.rPr)
        if source._r.rPr is not None:
            target._r.insert(0, deepcopy(source._r.rPr))

    @staticmethod
    def _set_language(run: Run, target_language: str) -> None:
        tag = _LANGUAGE_TAGS.get(target_language)
        if tag is None:
            return
        rPr = run._r.get_or_add_rPr()
        lang = rPr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rPr.insert_element_before(
                lang, "w:eastAsianLayout", "w:specVanish", "w:oMath"
            )
        if target_language in _EAST_ASIAN:
            lang.set(qn("w:eastAsia"), tag)
            rPr.get_or_add_rFonts().set(qn("w:hint"), "eastAsia")
        else:
            lang.set(qn("w:val"), tag)

    @staticmethod
    def _let_row_grow(paragraph: Paragraph) -> None:
        """A table row with an exact height would clip a longer translation."""
        for tr in paragraph._p.iterancestors(qn("w:tr")):
            for height in tr.iter(qn("w:trHeight")):
                if height.get(qn("w:hRule")) == "exact":
                    height.set(qn("w:hRule"), "atLeast")

    @staticmethod
    def _sync_fallback_text_boxes(element) -> None:
        """Copy translated text boxes into their legacy (VML) fallback, shown by older readers."""
        for alternate in element.iter(f"{_MC}AlternateContent"):
            choice = [
                box
                for branch in alternate.iterchildren(f"{_MC}Choice")
                for box in branch.iter(qn("w:txbxContent"))
            ]
            fallback = [
                box
                for branch in alternate.iterchildren(f"{_MC}Fallback")
                for box in branch.iter(qn("w:txbxContent"))
            ]
            if len(choice) != len(fallback):
                continue
            for source, target in zip(choice, fallback, strict=True):
                for child in list(target):
                    target.remove(child)
                target.extend(deepcopy(child) for child in source)
