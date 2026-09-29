"""Tests for writing the blocks added in the editor into a DOCX."""

import asyncio
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from pydantic import ValidationError

from ai_translation.domain.document import DOCXHandler
from ai_translation.infrastructure.rest.schemas import UpdateSegmentsRequest


def _sample_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Paragraf pertama")
    doc.add_paragraph("Paragraf kedua")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _png() -> bytes:
    output = BytesIO()
    Image.new("RGB", (40, 20), "red").save(output, format="PNG")
    return output.getvalue()


def _rebuild(insertions: list[dict], media: dict[str, bytes] | None = None) -> Document:
    handler = DOCXHandler()
    original = _sample_docx()
    translations = {"para_0": "First paragraph", "para_1": "Second paragraph"}
    content = asyncio.run(
        handler.create_translated_document(original, translations, "id", "en", insertions=insertions, media=media)
    )
    return Document(BytesIO(content))


def _body_texts(doc: Document) -> list[str]:
    texts = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            texts.append("".join(t.text or "" for t in child.iter(qn("w:t"))))
        elif child.tag == qn("w:tbl"):
            texts.append("<table>")
    return texts


def test_layout_gives_body_index():
    layout = DOCXHandler().document_layout(_sample_docx())
    assert [block["body_index"] for block in layout["sections"][0]["blocks"]] == [0, 1]


def test_blocks_are_placed_after_their_anchor_in_order():
    doc = _rebuild(
        [
            {"type": "paragraph", "after": -1, "kind": "h1", "runs": [{"text": "Title"}]},
            {"type": "paragraph", "after": 0, "kind": "normal", "runs": [{"text": "A"}]},
            {"type": "paragraph", "after": 0, "kind": "normal", "runs": [{"text": "B", "style": {"bold": True}}]},
            {"type": "divider", "after": 1},
        ]
    )
    assert _body_texts(doc) == ["Title", "First paragraph", "A", "B", "Second paragraph", ""]
    assert doc.paragraphs[0].style.style_id == "Heading1"
    assert doc.paragraphs[3].runs[0].bold is True
    # The translations are still where the original paragraphs were.
    assert doc.paragraphs[1].text == "First paragraph"


def test_lists_restart_numbering_per_run():
    doc = _rebuild(
        [
            {"type": "paragraph", "after": 0, "kind": "numbered", "runs": [{"text": "one"}]},
            {"type": "paragraph", "after": 0, "kind": "numbered", "level": 1, "runs": [{"text": "nested"}]},
            {"type": "paragraph", "after": 1, "kind": "numbered", "runs": [{"text": "again"}]},
            {"type": "paragraph", "after": 1, "kind": "bullet", "runs": [{"text": "dot"}]},
        ]
    )
    num_ids = {}
    for paragraph in doc.paragraphs:
        num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
        if num_pr is not None:
            num_ids[paragraph.text] = (num_pr.numId.val, num_pr.ilvl.val)
    assert num_ids["one"][0] == num_ids["nested"][0]
    assert num_ids["nested"][1] == 1
    # A list after other content starts again at 1: a new num.
    assert num_ids["again"][0] != num_ids["one"][0]
    assert num_ids["dot"][0] not in (num_ids["one"][0], num_ids["again"][0])


def test_table_with_header_row_and_line_breaks():
    doc = _rebuild(
        [
            {
                "type": "table",
                "after": 0,
                "header_row": True,
                "rows": [
                    {"cells": [{"paragraphs": [[{"text": "Name"}]]}, {"paragraphs": [[{"text": "Qty"}]]}]},
                    {"cells": [{"paragraphs": [[{"text": "Apple\nRed"}]]}, {"paragraphs": [[]]}]},
                ],
            }
        ]
    )
    assert _body_texts(doc) == ["First paragraph", "<table>", "Second paragraph"]
    table = doc.tables[0]
    assert table.cell(0, 0).text == "Name"
    assert table.cell(0, 0).paragraphs[0].runs[0].bold is True
    assert table.cell(1, 0).text == "Apple\nRed"
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None


def test_image_is_embedded():
    name = "upload/" + "a" * 32 + ".png"
    doc = _rebuild(
        [{"type": "image", "after": 1, "src": name, "width": 120, "height": 60, "align": "right", "alt": "Logo"}],
        media={name: _png()},
    )
    assert len(doc.inline_shapes) == 1
    assert doc.inline_shapes[0].width.pt == pytest.approx(120, abs=0.1)
    assert next(doc.element.body.iter(qn("wp:docPr"))).get("descr") == "Logo"


def test_missing_image_is_skipped():
    doc = _rebuild([{"type": "image", "after": 1, "src": "upload/" + "b" * 32 + ".png", "width": 1, "height": 1}])
    assert len(doc.inline_shapes) == 0


def test_segment_keys_are_unchanged_by_insertions():
    """Keys come from the original upload, which never holds the inserted paragraphs."""
    doc = _rebuild([{"type": "paragraph", "after": -1, "kind": "normal", "runs": [{"text": "Inserted"}]}])
    output = BytesIO()
    doc.save(output)
    rebuilt_keys = asyncio.run(DOCXHandler().extract_text(output.getvalue()))
    original_keys = asyncio.run(DOCXHandler().extract_text(_sample_docx()))
    assert set(original_keys) == {"para_0", "para_1"}
    # The rebuilt file has one more paragraph; the editor never re-reads it.
    assert len(rebuilt_keys) == 3


def test_request_validates_blocks():
    request = UpdateSegmentsRequest.model_validate(
        {"segments": [], "insertions": [{"type": "divider", "after": 2}, {"type": "paragraph", "after": 0}]}
    )
    assert request.insertions[1].kind == "normal"
    with pytest.raises(ValidationError):
        UpdateSegmentsRequest.model_validate(
            {"segments": [], "insertions": [{"type": "image", "after": 0, "src": "../etc", "width": 1, "height": 1}]}
        )
    with pytest.raises(ValidationError):
        UpdateSegmentsRequest.model_validate({"segments": [], "insertions": [{"type": "video", "after": 0}]})
