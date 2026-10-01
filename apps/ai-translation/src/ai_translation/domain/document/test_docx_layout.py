"""Tests for the editor's DOCX layout and for writing formatted edits back."""

import asyncio
from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.shared import Pt
from pydantic import ValidationError

from ai_translation.domain.document import DOCXHandler, PDFHandler
from ai_translation.infrastructure.rest.schemas import SegmentEdit


def _sample_docx() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Laporan Rahasia"
    doc.add_heading("Ringkasan", level=1)
    total = doc.add_paragraph("Pendapatan:\t$120,000")
    total.paragraph_format.tab_stops.add_tab_stop(Pt(400), WD_TAB_ALIGNMENT.RIGHT)
    doc.add_paragraph("Langkah pertama", style="List Number")
    doc.add_paragraph("Langkah kedua", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Bulan"
    table.cell(0, 1).text = "Keterangan"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Awal tahun"
    doc.add_paragraph("Akhir halaman").add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Halaman berikutnya")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _paragraphs(blocks: list[dict]):
    for block in blocks:
        if block["type"] == "paragraph":
            yield block
        elif block["type"] == "table":
            for row in block["rows"]:
                for cell in row["cells"]:
                    yield from _paragraphs(cell["blocks"])
        elif block["type"] == "frame":
            yield from _paragraphs(block["blocks"])


def _layout_paragraphs(layout: dict) -> list[dict]:
    blocks = [block for section in layout["sections"] for block in section["blocks"]]
    blocks += [block for header in layout["headers"].values() for block in header]
    return list(_paragraphs(blocks))


def test_layout_places_every_segment_once():
    content = _sample_docx()
    handler = DOCXHandler()
    keys = asyncio.run(handler.extract_text(content))
    layout = handler.document_layout(content)

    placed = [
        part["key"]
        for paragraph in _layout_paragraphs(layout)
        for part in paragraph["parts"]
        if part["type"] == "segment"
    ]
    assert sorted(placed) == sorted(keys)


def test_layout_keeps_structure_and_formatting():
    layout = DOCXHandler().document_layout(_sample_docx())
    section = layout["sections"][0]
    blocks = section["blocks"]
    paragraphs = [block for block in blocks if block["type"] == "paragraph"]

    assert section["page"]["width"] > 0 and section["header"]["default"] in layout["headers"]
    assert paragraphs[0]["heading"] == "h1"

    # "Label<tab>number": the label is translated, the number stays as it is.
    total = paragraphs[1]
    assert [part["type"] for part in total["parts"]] == ["segment", "tab", "text"]
    assert total["parts"][2]["text"] == "$120,000"
    assert any(tab["align"] == "right" and tab["pos"] == 400 for tab in total["style"]["tabs"])

    assert [p["marker"]["text"] for p in paragraphs[2:4]] == ["1.", "2."]

    table = next(block for block in blocks if block["type"] == "table")
    assert [len(row["cells"]) for row in table["rows"]] == [2, 1]
    assert table["rows"][1]["cells"][0]["span"] == 2
    assert table["rows"][0]["cells"][0]["borders"]["top"] is not None

    page_end = next(p for p in paragraphs if p["break_after"])
    assert page_end["parts"][0]["type"] == "segment"

    # Run styles are sent once and referred to by index.
    assert all(isinstance(part["run"], int) for p in paragraphs for part in p["parts"] if "run" in part)
    assert layout["run_styles"][paragraphs[0]["parts"][0]["run"]]["bold"] is True


def test_docx_export_applies_runs_and_paragraph_style():
    content = _sample_docx()
    handler = DOCXHandler()
    keys = asyncio.run(handler.extract_text(content))
    translations = {key: " ".join(texts) for key, texts in keys.items()}
    heading_key, body_key = "para_0", "para_1_0"
    translations[body_key] = "Revenue:"
    runs = {body_key: [{"text": "Rev", "style": {"bold": True}}, {"text": "enue:", "style": {"italic": True, "highlight": "#ffff00"}}]}
    styles = {heading_key: {"heading": "h2"}, body_key: {"align": "center", "font_size": 13}}

    output = asyncio.run(
        handler.create_translated_document(content, translations, "id", "en", styles=styles, runs=runs)
    )
    doc = Document(BytesIO(output))
    heading, body = doc.paragraphs[0], doc.paragraphs[1]

    assert heading.style.name == "Heading 2"
    written = [(run.text, run.bold, run.italic) for run in body.runs if run.text and run.text != "\t"]
    assert written[:2] == [("Rev", True, None), ("enue:", None, True)]
    assert body.alignment is not None and body.runs[0].font.size == Pt(13)
    # The fixed number after the tab is untouched.
    assert body.text.endswith("$120,000")


def test_pdf_export_writes_formatted_runs():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 100), "Revenue grew strongly this month.", fontsize=12)
    content = pdf.tobytes()
    handler = PDFHandler()
    key = next(iter(asyncio.run(handler.extract_text(content))))
    text = "Pendapatan tumbuh kuat."
    runs = {key: [{"text": "Pendapatan", "style": {"bold": True, "color": "#c00000"}}, {"text": " tumbuh kuat.", "style": {}}]}

    output = asyncio.run(handler.create_translated_document(content, {key: text}, "en", "id", runs=runs))
    with pymupdf.open(stream=output, filetype="pdf") as result:
        spans = [
            span
            for block in result[0].get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line["spans"]
        ]
    first = next(span for span in spans if span["text"].startswith("Pendapatan"))
    assert "Bold" in first["font"] and first["color"] == 0xC00000


def test_pdf_separates_company_name_from_right_aligned_date():
    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_htmlbox(
        pymupdf.Rect(50, 100, 580, 150),
        '<a href="https://labamu.example"><span style="font-weight:bold;color:#0040cc">'
        "Labamu MRP (Venture By Standard Chartered)"
        "</span></a>"
        + "&nbsp;" * 20
        + '<span style="font-style:italic;color:#000000">'
        "17 Juni 2026 - 30 September 2026</span>",
        css="* {font-family: Times; font-size: 12px; margin: 0; padding: 0;}",
    )
    page.insert_link(
        {
            "kind": pymupdf.LINK_URI,
            "from": pymupdf.Rect(51, 99, 302, 117),
            "uri": "https://labamu.example",
        }
    )
    page.insert_text((540, 300), "x", fontsize=12)

    handler = PDFHandler()
    content = pdf.tobytes()
    with pymupdf.open(stream=content, filetype="pdf") as source:
        page = source[0]
        segments, layout = handler._page_segments(page)

        assert [segment.text for segment in segments[:2]] == [
            "Labamu MRP (Venture By Standard Chartered)",
            "17 Juni 2026 - 30 September 2026",
        ]
        company, date = segments[:2]
        assert company.style.bold and not company.style.italic
        assert date.style.italic and not date.style.bold
        assert company.style.color != date.style.color
        assert date.bbox.x0 - company.bbox.x1 > 50
        assert layout.text_right > date.bbox.x1 + date.style.font_size * 0.35
        assert handler._place(date, layout).align == "right"
        assert max(link["from"].x1 for link in page.get_links()) < date.bbox.x0

    translations = {segment.key: segment.text for segment in segments}
    output = asyncio.run(
        handler.create_translated_document(content, translations, "en", "id")
    )
    with pymupdf.open(stream=output, filetype="pdf") as translated:
        page = translated[0]
        spans = [
            span
            for block in page.get_text("dict")["blocks"]
            if block.get("type") == 0
            for line in block["lines"]
            for span in line["spans"]
        ]
        company_span = next(span for span in spans if span["text"].startswith("Labamu"))
        date_span = next(span for span in spans if span["text"].startswith("17 Juni"))
        assert "Bold" in company_span["font"] and company_span["color"] == 0x0040CC
        assert "Italic" in date_span["font"] and date_span["color"] == 0
        assert max(link["from"].x1 for link in page.get_links()) < date_span["bbox"][0]


def test_segment_edit_runs_must_match_text():
    SegmentEdit(key="para_1", translated_text="ab", runs=[{"text": "a"}, {"text": "b", "style": {"bold": True}}])
    with pytest.raises(ValidationError):
        SegmentEdit(key="para_1", translated_text="ab", runs=[{"text": "a"}])
