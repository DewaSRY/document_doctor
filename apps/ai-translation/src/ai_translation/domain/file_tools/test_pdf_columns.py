"""Tests that side-by-side regions of a PDF page stay side by side in every Word reader."""

from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Twips

from ai_translation.domain.file_tools import pdf_to_docx


def _picture(width: int, height: int) -> bytes:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), 0)
    pixmap.set_rect(pixmap.irect, (200, 60, 60))
    return pixmap.tobytes("png")


def _text_beside_photo(page: pymupdf.Page, top: float) -> None:
    """A heading, ten lines of text with a photo to their right, then a paragraph below."""
    page.insert_text((72, top), "Heading above the columns", fontsize=14)
    for number in range(10):
        page.insert_text((72, top + 40 + 14 * number), f"Line {number} of the text beside a photo.", fontsize=10)
    page.insert_image(pymupdf.Rect(330, top + 28, 520, top + 170), stream=_picture(40, 30))
    page.insert_text((72, top + 220), "Paragraph below the columns, across the width of the page.", fontsize=10)


def _convert(pdf: pymupdf.Document):
    return Document(BytesIO(pdf_to_docx(pdf.tobytes())))


def _text(element) -> str:
    return "".join(t.text or "" for t in element.iter(qn("w:t")))


def test_text_beside_a_photo_becomes_a_borderless_table_of_two_cells():
    pdf = pymupdf.open()
    _text_beside_photo(pdf.new_page(), 80)

    doc = _convert(pdf)
    body = doc.element.body

    # No column sections are left for readers other than Word to drop.
    assert not body.xpath('.//w:type[@w:val="nextColumn"]')
    assert not body.xpath(".//w:cols[@w:num]")
    assert len(doc.tables) == 1
    left, right = doc.tables[0]._tbl.tr_lst[0].tc_lst
    assert "Line 0 of the text beside a photo." in _text(left)
    assert "Line 9 of the text beside a photo." in _text(left)
    assert right.xpath(".//w:drawing")
    assert not left.xpath(".//w:drawing")
    assert not doc.tables[0]._tbl.xpath("./w:tblPr/w:tblBorders/*[@w:val!='nil']")
    # The row keeps its columns on one page.
    assert doc.tables[0]._tbl.xpath("./w:tr/w:trPr/w:cantSplit")
    # The text before and after the region stays outside the table.
    assert "Heading above the columns" not in _text(doc.tables[0]._tbl)
    assert "Paragraph below the columns" not in _text(doc.tables[0]._tbl)
    assert "Paragraph below the columns" in _text(body)


def test_table_columns_are_as_wide_as_the_columns_on_the_page():
    pdf = pymupdf.open()
    _text_beside_photo(pdf.new_page(), 80)

    doc = _convert(pdf)

    grid = [int(column.get(qn("w:w"))) for column in doc.tables[0]._tbl.tblGrid.iter(qn("w:gridCol"))]
    section = doc.sections[0]
    assert section.page_width is not None
    assert section.left_margin is not None
    assert section.right_margin is not None
    assert Twips(sum(grid)).pt == pytest.approx((section.page_width - section.left_margin - section.right_margin) / 12700)
    # The first column ends between the end of the text and the photo, at x=330.
    text_end = 72 + pymupdf.get_text_length("Line 0 of the text beside a photo.", fontsize=10)
    assert text_end < section.left_margin.pt + Twips(grid[0]).pt <= 330


def test_region_at_the_bottom_of_a_page_is_rewritten_too():
    # pdf2docx ends a page's last column only when it starts the next page.
    pdf = pymupdf.open()
    _text_beside_photo(pdf.new_page(), 560)
    pdf.new_page().insert_text((72, 100), "Second page text.", fontsize=11)
    _text_beside_photo(pdf.new_page(), 80)

    doc = _convert(pdf)
    body = doc.element.body

    assert len(doc.tables) == 2
    assert all(table._tbl.tr_lst[0].tc_lst[1].xpath(".//w:drawing") for table in doc.tables)
    assert not body.xpath('.//w:type[@w:val="nextColumn"]')
    assert not body.xpath(".//w:cols[@w:num]")
    assert len(doc.sections) == 3


def test_section_breaks_left_inside_a_page_are_joined():
    pdf = pymupdf.open()
    _text_beside_photo(pdf.new_page(), 80)

    doc = _convert(pdf)

    # Continuous breaks between single-column parts of a page separate nothing.
    assert len(doc.sections) == 1
    assert not doc.element.body.xpath('.//w:p/w:pPr/w:sectPr/w:type[@w:val="continuous"]')
