"""
Carry a PDF's standalone lines over to the Word document pdf2docx rebuilds.

pdf2docx turns lines into table borders and into underlines or strikes of the
text they touch, and drops every other line: rules under headings, separators
between paragraphs, fill-in lines of forms. This module collects those lines
from the parsed layout and draws each one as a shape behind the text, at the
position, length, thickness and color it has on the PDF page.
"""

from collections.abc import Iterator
from dataclasses import dataclass

import pymupdf
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

# A shape is a line when it is at most this thick, in points...
_MAX_THICKNESS = 6.0
# ...and at least this many times as long as it is thick.
_MIN_ASPECT = 4.0
# Lines thinner than this would not show; draw them this thick instead.
_MIN_THICKNESS = 0.25

_SHAPE_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_WPS = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'


@dataclass(frozen=True)
class PdfLine:
    """A horizontal or vertical line of a page, as the rectangle it covers."""

    rect: pymupdf.Rect
    color: int  # 0xRRGGBB


def page_lines(page) -> list[PdfLine]:
    """The lines of a parsed pdf2docx page that it will not write itself."""
    lines: dict[tuple, PdfLine] = {}
    for shape in _shapes(page):
        # Shapes pdf2docx uses, as table borders or text styles, have a type.
        if shape.is_determined:
            continue
        rect = pymupdf.Rect(shape.bbox)
        thickness, length = sorted((rect.width, rect.height))
        if thickness > _MAX_THICKNESS or length < _MIN_ASPECT * max(thickness, _MIN_THICKNESS):
            continue
        # A table's cells are given the shapes around them too.
        key = tuple(round(v, 1) for v in (rect.x0, rect.y0, rect.x1, rect.y1))
        lines.setdefault(key, PdfLine(rect, shape.color or 0))
    return list(lines.values())


def _shapes(page) -> Iterator:
    for section in page.sections:
        for column in section:
            yield from column.shapes
            yield from _cell_shapes(column.blocks)


def _cell_shapes(blocks) -> Iterator:
    for block in blocks:
        if block.is_table_block:
            for row in block:
                for cell in row:
                    if cell is not None:
                        yield from cell.shapes
                        yield from _cell_shapes(cell.blocks)


def add_lines(paragraph: Paragraph, lines: list[PdfLine]) -> None:
    """Draw the lines of a page, anchored to its first paragraph and placed on the page."""
    # Inside a table, Word would place the shape relative to the cell.
    in_cell = any(True for _ in paragraph._p.iterancestors(qn("w:tc")))
    for line in lines:
        run = paragraph.add_run()
        run._r.append(_line_drawing(line, run.part.next_id, in_cell))


def _line_drawing(line: PdfLine, shape_id: int, in_cell: bool):
    rect = pymupdf.Rect(line.rect)
    if rect.height < _MIN_THICKNESS:
        middle = (rect.y0 + rect.y1) / 2
        rect.y0, rect.y1 = middle - _MIN_THICKNESS / 2, middle + _MIN_THICKNESS / 2
    if rect.width < _MIN_THICKNESS:
        middle = (rect.x0 + rect.x1) / 2
        rect.x0, rect.x1 = middle - _MIN_THICKNESS / 2, middle + _MIN_THICKNESS / 2
    x, y, cx, cy = (int(Pt(v)) for v in (rect.x0, rect.y0, rect.width, rect.height))
    return parse_xml(
        f'<w:drawing {nsdecls("w", "wp", "a")} {_WPS}>'
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{shape_id}"'
        f' behindDoc="1" locked="0" layoutInCell="{0 if in_cell else 1}" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{x}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{y}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        "<wp:wrapNone/>"
        f'<wp:docPr id="{shape_id}" name="Line {shape_id}"/>'
        "<wp:cNvGraphicFramePr/>"
        f'<a:graphic><a:graphicData uri="{_SHAPE_URI}">'
        "<wps:wsp><wps:cNvSpPr/>"
        "<wps:spPr>"
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:solidFill><a:srgbClr val="{line.color:06X}"/></a:solidFill>'
        "<a:ln><a:noFill/></a:ln>"
        "</wps:spPr>"
        '<wps:bodyPr lIns="0" tIns="0" rIns="0" bIns="0"/>'
        "</wps:wsp>"
        "</a:graphicData></a:graphic>"
        "</wp:anchor></w:drawing>"
    )
